# Шаг 19: таблица аналитики (xlsx, формулы)
import csv, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
M='/home/user/CreateVideo/deliverables/yafho-silicare-instagram/MASTER'
F='Arial'; f=lambda **k: Font(name=F,size=k.pop('size',10),**k)
HDR=PatternFill('solid',fgColor='2B2B2B'); INP=PatternFill('solid',fgColor='FFF2CC'); OR=PatternFill('solid',fgColor='F38221')
GOOD=PatternFill('solid',fgColor='D9EAD3'); BAD=PatternFill('solid',fgColor='F4CCCC')
thin=Border(bottom=Side(style='thin',color='DDDDDD'))
def header(ws,row,cols,widths=None):
    for i,c in enumerate(cols,1):
        x=ws.cell(row,i,c); x.font=f(bold=True,color='FFFFFF'); x.fill=HDR; x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
        if widths: ws.column_dimensions[CL(i)].width=widths[i-1]
    ws.row_dimensions[row].height=42; ws.freeze_panes=ws.cell(row+1,1)
wb=Workbook()
# --- Инструкция
ws=wb.active; ws.title='Инструкция'; ws.column_dimensions['A'].width=110
lines=[('Аналитика Instagram · Yafho «Кожа помнит»',dict(bold=True,size=14)),('',{}),
 ('Как заполнять',dict(bold=True)),
 ('1. Раз в неделю (понедельник) откройте статистику каждой публикации прошлой недели и внесите цифры в ЖЁЛТЫЕ ячейки листа «Публикации».',{}),
 ('2. Всё белое считается само: доли, флаги «цель выполнена», сводки на листах «Недели» и «Рубрики». Формулы не трогайте.',{}),
 ('3. Цели (пороги) — на листе «Цели». Поменяйте там — флаги пересчитаются.',{}),
 ('4. Строка 12.10 заполнена ПРИМЕРОМ (серый курсив) — перезапишите её реальными цифрами.',{}),
 ('5. «A/B обложки»: вносите переходы из профиля и показы за каждые 72 ч с обложкой — победитель определяется автоматически.',{}),
 ('6. «Реклама»: расход и результаты из Ads Manager. «Заявки»: одна строка на диалог в директе (код КАТАЛОГ / ЛУПА).',{}),
 ('',{}),('Где брать цифры в Instagram',dict(bold=True)),
 ('Охват, просмотры, досмотр 3 с (Reels: «Удержание»), среднее время просмотра, лайки, комментарии, сохранения, репосты, переходы в профиль, подписки — «Статистика» под публикацией.',{}),
 ('',{}),('Обозначения: жёлтый фон — вводить; зелёный/красный — цель выполнена / нет.',dict(italic=True,color='777777'))]
for i,(t,st) in enumerate(lines,1): c=ws.cell(i,1,t); c.font=f(**st); c.alignment=Alignment(wrap_text=True)
# --- Цели
wc=wb.create_sheet('Цели'); header(wc,1,['Метрика','Цель','Комментарий'],[34,12,70])
G=[('Досмотр первых 3 с (Reels)',0.6,'Хук работает, если > 60% досматривают 3 с'),('Доля репостов (репосты/охват)',0.01,'> 1% — контент «пересылают»'),
   ('Доля сохранений (сохранения/охват)',0.02,'> 2% — полезный контент (карусели)'),('Вовлечённость ER (реакции/охват)',0.05,'лайки+комментарии+сохранения+репосты'),
   ('Подписки на 1000 охвата',3,'сколько подписок приносит 1000 охвата')]
for i,(a,b,c) in enumerate(G,2):
    wc.cell(i,1,a).font=f(); x=wc.cell(i,2,b); x.font=f(color='0000FF'); x.fill=INP; x.number_format='0.0%' if b<1 else '0'
    wc.cell(i,3,c).font=f(color='555555')
wc.cell(8,1,'Источник целей: ориентиры команды (MASTER_CALENDAR.md → «Метрики по неделям»), а не отраслевой стандарт. Уточните после первого месяца.').font=f(italic=True,color='777777')
# --- Публикации
wp=wb.create_sheet('Публикации')
cols=['Дата','Нед.','Формат','Рубрика','Публикация','Файл','Охват','Просмотры','Досмотр 3 с, %','Ср. время просмотра, с','Лайки','Комментарии','Сохранения','Репосты','Переходы в профиль','Подписки','Сообщения (код)',
      'ER','Доля сохранений','Доля репостов','Подписки / 1000 охвата','Досмотр ≥ цели?','Репосты ≥ цели?','Сохранения ≥ цели?']
header(wp,1,cols,[11,6,10,16,42,28,10,11,10,11,8,9,10,9,10,9,10,8,10,10,10,10,10,10])
rows=list(csv.DictReader(open(f'{M}/master_calendar.csv')))
wk={t:i+1 for i,t in enumerate(dict.fromkeys(r['week'] for r in rows))}
n=len(rows)
for i,r in enumerate(rows,2):
    w=wk[r['week']]; w='после' if w>7 else w
    vals=[dt.date.fromisoformat(r['date']),w,r['format'],r['rubric'],r['title'],r['file']]
    for j,v in enumerate(vals,1):
        c=wp.cell(i,j,v); c.font=f(); c.border=thin
    wp.cell(i,1).number_format='DD.MM.YYYY'
    for j in range(7,18):
        c=wp.cell(i,j); c.fill=INP; c.font=f(color='0000FF'); c.border=thin
        c.number_format='0.0%' if j==9 else '#,##0'
    R=i
    fx={18:f'=IF(G{R}>0,(K{R}+L{R}+M{R}+N{R})/G{R},"")',19:f'=IF(G{R}>0,M{R}/G{R},"")',20:f'=IF(G{R}>0,N{R}/G{R},"")',21:f'=IF(G{R}>0,P{R}/G{R}*1000,"")',
        22:f'=IF(OR(C{R}<>"Reel",I{R}=""),"",IF(I{R}>=Цели!$B$2,"да","нет"))',23:f'=IF(T{R}="","",IF(T{R}>=Цели!$B$3,"да","нет"))',24:f'=IF(S{R}="","",IF(S{R}>=Цели!$B$4,"да","нет"))'}
    for j,v in fx.items():
        c=wp.cell(i,j,v); c.font=f(); c.border=thin; c.number_format='0.0' if j==21 else '0.0%'
# пример в первой строке
ex=[18400,26100,0.64,9.2,812,47,265,231,390,61,4]
for j,v in enumerate(ex,7):
    c=wp.cell(2,j,v); c.font=f(color='999999',italic=True)
wp.cell(2,7).comment=Comment('ПРИМЕР — перезапишите реальными цифрами из статистики','Yafho')
last=n+1
for col in 'VWX':
    rng=f'{col}2:{col}{last}'
    wp.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"да"'],fill=GOOD))
    wp.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"нет"'],fill=BAD))
wp.auto_filter.ref=f'A1:X{last}'
P=f"Публикации!"
# --- Недели
ww=wb.create_sheet('Недели'); header(ww,1,['Неделя','Тема','Публикаций','С цифрами','Охват','Сохранения','Репосты','Подписки','Сообщения','Средний ER','Ср. досмотр 3 с (Reels)'],[8,46,11,10,11,11,10,10,11,11,12])
titles=list(dict.fromkeys(r['week'] for r in rows))
for i,t in enumerate(titles[:7],2):
    k=i-1; ww.cell(i,1,k); ww.cell(i,2,t)
    A=f'{P}$B$2:$B${last}'
    fx=[f'=COUNTIF({A},A{i})',f'=COUNTIFS({A},A{i},{P}$G$2:$G${last},">0")',f'=SUMIFS({P}$G$2:$G${last},{A},A{i})',f'=SUMIFS({P}$M$2:$M${last},{A},A{i})',
        f'=SUMIFS({P}$N$2:$N${last},{A},A{i})',f'=SUMIFS({P}$P$2:$P${last},{A},A{i})',f'=SUMIFS({P}$Q$2:$Q${last},{A},A{i})',
        f'=IFERROR(AVERAGEIFS({P}$R$2:$R${last},{A},A{i},{P}$G$2:$G${last},">0"),"")',f'=IFERROR(AVERAGEIFS({P}$I$2:$I${last},{A},A{i},{P}$C$2:$C${last},"Reel",{P}$I$2:$I${last},">0"),"")']
    for j,v in enumerate(fx,3):
        c=ww.cell(i,j,v); c.number_format='0.0%' if j>=10 else '#,##0'
    for j in range(1,12): ww.cell(i,j).font=f(); ww.cell(i,j).border=thin
r=len(titles[:7])+2; ww.cell(r,2,'Итого').font=f(bold=True)
for j in range(3,10): c=ww.cell(r,j,f'=SUM({CL(j)}2:{CL(j)}{r-1})'); c.font=f(bold=True); c.number_format='#,##0'
# --- Рубрики
wr=wb.create_sheet('Рубрики'); header(wr,1,['Рубрика','Публикаций','Охват','Средний ER','Доля сохранений','Доля репостов','Подписки','Сообщения'],[20,11,11,11,12,12,10,11])
rubs=sorted(set(r['rubric'] for r in rows))
for i,rb in enumerate(rubs,2):
    wr.cell(i,1,rb); A=f'{P}$D$2:$D${last}'; G_=f'{P}$G$2:$G${last}'
    fx=[f'=COUNTIF({A},A{i})',f'=SUMIFS({G_},{A},A{i})',f'=IFERROR(AVERAGEIFS({P}$R$2:$R${last},{A},A{i},{G_},">0"),"")',
        f'=IFERROR(SUMIFS({P}$M$2:$M${last},{A},A{i})/C{i},"")',f'=IFERROR(SUMIFS({P}$N$2:$N${last},{A},A{i})/C{i},"")',f'=SUMIFS({P}$P$2:$P${last},{A},A{i})',f'=SUMIFS({P}$Q$2:$Q${last},{A},A{i})']
    for j,v in enumerate(fx,2):
        c=wr.cell(i,j,v); c.number_format='0.0%' if j in (4,5,6) else '#,##0'
    for j in range(1,9): wr.cell(i,j).font=f(); wr.cell(i,j).border=thin
# --- A/B обложки
wa=wb.create_sheet('A-B обложки'); header(wa,1,['Ролик','Показы профиля A','Переходы A','Показы профиля B','Переходы B','Показы профиля C','Переходы C','CTR A','CTR B','CTR C','Победитель'],[22,12,10,12,10,12,10,9,9,9,11])
AB=['m1_manifest','m2_pismo','m3_bumaga','m4_persik','m5_lepestok','e1_kesarevo','e2_oritn','e3_shram','e4_babochki','e5_medsestra']
for i,k in enumerate(AB,2):
    wa.cell(i,1,k).font=f()
    for j in range(2,8): c=wa.cell(i,j); c.fill=INP; c.font=f(color='0000FF'); c.number_format='#,##0'
    for j,(a,b) in enumerate((('B','C'),('D','E'),('F','G')),8):
        c=wa.cell(i,j,f'=IF({a}{i}>0,{b}{i}/{a}{i},"")'); c.number_format='0.0%'; c.font=f()
    c=wa.cell(i,11,f'=IF(COUNT(H{i}:J{i})<3,"идёт тест",INDEX({{"A","B","C"}},MATCH(MAX(H{i}:J{i}),H{i}:J{i},0)))'); c.font=f(bold=True)
r=len(AB)+3; wa.cell(r,1,'Итог по типам:').font=f(bold=True)
for j,(lbl,col) in enumerate((('A — вопрос','K'),('B — факт','K'),('C — эмоция','K'))):
    wa.cell(r+1+j,1,lbl).font=f(); wa.cell(r+1+j,2,f'=COUNTIF(K2:K{len(AB)+1},"{lbl[0]}")').font=f()
wa.cell(r+5,1,'Тип-победитель → на обложки следующих роликов (см. v10_ab_covers/README.md).').font=f(italic=True,color='777777')
# --- Реклама
wd=wb.create_sheet('Реклама'); header(wd,1,['Кампания','Ролик','Цель','Расход','Показы','Охват','Результаты','Цена результата','CPM','Частота'],[30,14,22,11,11,11,11,13,10,9])
C=[('A. Охват — «Персик»','m4','ThruPlay'),('B. Знакомство — «Манифест»','m1','подписки'),('C. Медсёстры — «Хирургия»','k5','сообщения'),('D. Клиники — «Линейка»','e6','квалиф. заявки')]
for i,(a,b,g) in enumerate(C,2):
    for j,v in enumerate((a,b,g),1): wd.cell(i,j,v).font=f()
    for j in range(4,8): c=wd.cell(i,j); c.fill=INP; c.font=f(color='0000FF'); c.number_format='#,##0.00' if j==4 else '#,##0'
    for j,v,nf in ((8,f'=IF(G{i}>0,D{i}/G{i},"")','#,##0.00'),(9,f'=IF(E{i}>0,D{i}/E{i}*1000,"")','#,##0.00'),(10,f'=IF(F{i}>0,E{i}/F{i},"")','0.0')):
        c=wd.cell(i,j,v); c.number_format=nf; c.font=f()
wd.cell(6,1,'Итого').font=f(bold=True)
for j in (4,5,7): c=wd.cell(6,j,f'=SUM({CL(j)}2:{CL(j)}5)'); c.font=f(bold=True); c.number_format='#,##0'
wd.cell(8,1,'Расход — в валюте рекламного кабинета. Частота > 3 или падение CTR на 30% → менять креатив (ADS.md).').font=f(italic=True,color='777777')
wd.conditional_formatting.add('J2:J5',CellIsRule(operator='greaterThan',formula=['3'],fill=BAD))
# --- Заявки
wl=wb.create_sheet('Заявки'); header(wl,1,['Дата','Instagram','Код','Из публикации','Тип','Отделение','Регион','Контакт','Статус','Следующий шаг','Ответственный'],[11,16,10,24,12,14,12,16,12,20,14])
wl.cell(2,1,dt.date(2026,11,2)).number_format='DD.MM.YYYY'
for j,v in enumerate(['@example','КАТАЛОГ','g1_katalog_za_30_sekund','клиника','ОРИТН','[город]','[телефон]','квалифицирована','звонок 04.11','[менеджер]'],2): wl.cell(2,j,v)
for j in range(1,12): wl.cell(2,j).font=f(color='999999',italic=True)
wl.cell(2,1).comment=Comment('ПРИМЕР — удалите или перезапишите','Yafho')
for col,opts in (('C','"КАТАЛОГ,ЛУПА,другое"'),('E','"клиника,аптека,дистрибьютор,пациент"'),('I','"новая,квалифицирована,образцы,заказ,отказ"')):
    dv=DataValidation(type='list',formula1=opts,allow_blank=True); wl.add_data_validation(dv); dv.add(f'{col}2:{col}500')
# воронка
wl.cell(1,13,'Воронка').font=f(bold=True,color='FFFFFF'); wl.cell(1,13).fill=HDR; wl.cell(1,14).fill=HDR; wl.cell(1,15).fill=HDR
wl.column_dimensions['M'].width=18; wl.cell(1,14,'КАТАЛОГ').font=f(bold=True,color='FFFFFF'); wl.cell(1,15,'ЛУПА').font=f(bold=True,color='FFFFFF')
st=['все заявки','квалифицирована','образцы','заказ']
for i,s in enumerate(st,2):
    wl.cell(i,13,s).font=f()
    for j,code in ((14,'КАТАЛОГ'),(15,'ЛУПА')):
        fm=f'=COUNTIF($C$2:$C$500,"{code}")' if s=='все заявки' else f'=COUNTIFS($C$2:$C$500,"{code}",$I$2:$I$500,"{s}")'
        wl.cell(i,j,fm).font=f()
wl.cell(7,13,'Каждая заявка считается только в своём текущем статусе (не накопительно).').font=f(italic=True,color='777777',size=9)
for s in wb.worksheets: s.sheet_view.showGridLines=False
wb.save(f'{M}/analytics_tracker.xlsx'); print('saved')
